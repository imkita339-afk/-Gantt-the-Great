"""Мастер таблиц, «Гант-таблица» (Excel/CSV) и MS Project — туда и обратно."""
import json

import openpyxl
import pytest

from gantt.exchange import check
from gantt.formats import NeedsWizard, open_path, save_template, table_result
from gantt.formats import gantt_table, msproject
from gantt.formats.table import guess_mapping, read_clipboard_text, read_csv, read_xlsx
from gantt.model import Dataset

CSV = """Код;Тема;Состояние;Исполнитель;Начало;Срок;Оценка;Приоритет
HD-1;Новый дизайн;В работе;Дизайнер;14.09.2026;30.09.2026;24;Высокий
HD-2;Оплата частями;Открыта;Разработчик;01.10.2026;20.10.2026;40;Низкий
HD-3;Ошибка 500;Готово;Разработчик, Тестировщик;22.09.2026;23.09.2026;;Средний
"""


def test_csv_flat_with_status_column(tmp_path):
    p = tmp_path / "zadachi.csv"
    p.write_text(CSV, encoding="utf-8-sig")
    with pytest.raises(NeedsWizard) as e:
        open_path(p)
    table = e.value.table
    m = guess_mapping(table)
    assert m.columns["number"] == 0 and m.columns["key"] == 0 and m.columns["name"] == 1 and m.columns["status"] == 2
    assert m.hierarchy == "flat" and m.status_mode == "column"
    assert m.status_map == {"В работе": "inProgress", "Открыта": "notStarted", "Готово": "done"}
    m.custom = {"Оценка": 6}
    m.template = "Задачи по сайту"
    res = table_result(table, m)
    ds = Dataset(res.document)
    hd3 = ds.nodes["HD-3"]
    assert ds.status_kind(hd3) == "done" and len(hd3.raw["executors"]) == 2
    assert ds.nodes["HD-1"].raw["custom"] and ds.scale_item(ds.nodes["HD-1"], "priority")["order"] == 4
    save_template(m)                      # шаблон — дальше без мастера
    again = open_path(p)
    assert len(again.document["nodes"]) == 3


def test_xlsx_outline_levels(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Наименование", "Начало", "Окончание", "Исполнитель"])
    rows = [(0, "Раздел А"), (1, "Проект 1"), (2, "Задача 1.1"), (2, "Задача 1.2"), (1, "Проект 2"),
            (2, "Задача 2.1"), (0, "Раздел Б"), (1, "Проект 3")]
    for i, (lvl, name) in enumerate(rows, start=2):
        ws.append([name, "01.09.2026", "30.09.2026", "Аналитик 1"])
        ws.row_dimensions[i].outline_level = lvl
    p = tmp_path / "plan.xlsx"
    wb.save(p)
    with pytest.raises(NeedsWizard) as e:
        open_path(p)
    m = guess_mapping(e.value.table)
    assert m.hierarchy == "outline"
    ds = Dataset(table_result(e.value.table, m).document)
    assert ds.counts() == {"group": 2, "project": 3, "task": 3}
    assert [n.name for n in ds.roots[0].children] == ["Проект 1", "Проект 2"]


def test_numbering_hierarchy_from_clipboard():
    text = "№\tНазвание\tСрок\n1\tЭтап\t\n1.1\tПодэтап\t01.10.2026\n1.2\tЕщё\t02.10.2026\n2\tВторой\t\n"
    table = read_clipboard_text(text)
    m = guess_mapping(table)
    assert m.hierarchy == "numbering"
    ds = Dataset(table_result(table, m).document)
    assert [len(n.children) for n in ds.roots] == [2, 0]


def test_gantt_table_xlsx_roundtrip(root, tmp_path):
    doc = json.loads((root / "examples" / "full.gantt.json").read_text(encoding="utf-8"))
    p = tmp_path / "snimok.xlsx"
    gantt_table.write_xlsx(doc, p, colors={doc["nodes"][1]["id"]: "#A7C4E2"}, chart=True)
    wb = openpyxl.load_workbook(p)
    assert {"Данные", "Справочники", "Диаграмма", "_gantt"} <= set(wb.sheetnames)
    back = open_path(p)
    assert back.source_label == "Excel · Гант-таблица"
    assert len(back.document["nodes"]) == len(doc["nodes"])
    a, b = Dataset(doc), Dataset(back.document)
    assert a.counts() == b.counts()
    n0 = doc["nodes"][5]
    m0 = next(n for n in back.document["nodes"] if n["id"] == n0["id"])
    for f in ("planStart", "planEnd", "factStart", "factEnd", "status", "executors", "parentId"):
        assert m0.get(f) == n0.get(f), f


def test_gantt_table_csv_roundtrip(root, tmp_path):
    doc = json.loads((root / "examples" / "full.gantt.json").read_text(encoding="utf-8"))
    p = tmp_path / "snimok.csv"
    gantt_table.write_csv(doc, p)
    back = open_path(p)
    assert back.source_label == "CSV · Гант-таблица"
    assert Dataset(back.document).counts() == Dataset(doc).counts()


def test_msproject_roundtrip(root, tmp_path):
    doc = json.loads((root / "examples" / "other-system-flat.gantt.json").read_text(encoding="utf-8"))
    p = tmp_path / "plan.xml"
    msproject.write(doc, p)
    back = open_path(p)
    assert check(back.document).errors == []
    ids = {n["id"] for n in back.document["nodes"]}
    assert ids == {n["id"] for n in doc["nodes"]}           # «ID Ганта» сохранился
    hd102 = next(n for n in back.document["nodes"] if n["id"] == "HD-102")
    assert hd102["predecessors"][0]["id"] == "HD-101"
    hd105 = next(n for n in back.document["nodes"] if n["id"] == "HD-105")
    assert hd105.get("milestone") is True


def test_changes_file_is_recognized(root):
    res = open_path(root / "examples" / "changes.gantt.json")
    assert res.kind == "changes" and len(res.document["changes"]) == 3
