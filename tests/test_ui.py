"""Окно целиком, без экрана: открыть, поправить, сохранить во все форматы, принять правки и ответ источника."""
import json
from datetime import date

import openpyxl
import pytest

from gantt.exchange import check


@pytest.fixture
def win(root, make_win):
    w = make_win()
    assert w.open_file(str(root / "examples" / "full.gantt.json"))
    return w


def first_task(w):
    return next(n for n in w.ds.walk() if w.ds.role(n) == "task" and n.source("planEnd"))


def test_edit_survives_json_snapshot_roundtrip(win, tmp_path):
    task = first_task(win)
    win.push(task, {"planEnd": date(2030, 1, 1)}, "тест")
    path = tmp_path / "снимок.gantt.json"
    assert win.write_as("json", str(path), {})
    doc = json.loads(path.read_text(encoding="utf-8"))
    assert [c["field"] for c in doc["changes"]] == ["planEnd"]
    # правка не теряется: открыли снимок — она на месте, исходный набор её тоже помнит
    win.open_file(str(path))
    assert win.ds.get(win.ds.nodes[task.id], "planEnd") == date(2030, 1, 1)
    assert win.store.load(win.ds.key)[(task.id, "planEnd")] == date(2030, 1, 1)


def test_save_all_formats_and_open_back(win, tmp_path):
    task = first_task(win)
    win.push(task, {"planEnd": date(2030, 1, 1)}, "тест")
    for key, name in (("xlsx", "a.xlsx"), ("csv", "a.csv"), ("msproject", "a.xml")):
        p = tmp_path / name
        assert win.write_as(key, str(p), {"chart": True}), key
        before = win.ds.key
        assert win.open_file(str(p)), key
        assert win.ds.key != before or key == "json"          # сохранённая таблица — свой набор данных
        assert win.ds.get(win.ds.nodes[task.id], "planEnd") == date(2030, 1, 1), key
        win.open_file(str(win_root(win) / "examples" / "full.gantt.json"))
        assert win.ds.get(win.ds.nodes[task.id], "planEnd") == date(2030, 1, 1)  # исходная правка цела
    wb = openpyxl.load_workbook(tmp_path / "a.xlsx")
    assert "Диаграмма" in wb.sheetnames


def win_root(w):
    from gantt.paths import resource_path
    return resource_path("examples", "full.gantt.json").parent.parent


def test_png_and_pdf(win, tmp_path):
    from PySide6.QtGui import QImage

    png, pdf = tmp_path / "d.png", tmp_path / "d.pdf"
    assert win.write_as("png", str(png), {"scope": "visible", "title": "Портфель"})
    img = QImage(str(png))
    assert img.width() > 800 and img.height() > 100
    assert win.write_as("pdf", str(pdf), {"scope": "all", "page": "A4", "title": "Портфель"})
    assert pdf.read_bytes()[:4] == b"%PDF" and pdf.stat().st_size > 5000


def test_change_list_and_changes_pack(win, tmp_path):
    task = first_task(win)
    win.push(task, {"planEnd": date(2030, 1, 1)}, "тест")
    win.push(task, {"color": "#F4B6C2"}, "тест")
    assert win.write_change_list(str(tmp_path / "list.xlsx"))
    ws = openpyxl.load_workbook(tmp_path / "list.xlsx").active
    assert ws["A4"].value == "Раздел / проект" and ws.max_row == 6          # две правки
    assert win.write_change_list(str(tmp_path / "list.pdf"))
    pack = win.changes_pack()
    assert check(pack).errors == [] and [c["field"] for c in pack["changes"]] == ["planEnd"]


def test_incoming_changes_and_results(win, monkeypatch):
    from gantt.ui import incoming

    task = first_task(win)
    win.push(task, {"planEnd": date(2030, 1, 1)}, "тест")
    pack = win.changes_pack()
    win.push(task, {"planEnd": task.source("planEnd")}, "возврат")
    assert not win.ds.edit_count
    monkeypatch.setattr(incoming.ChangesImportDialog, "exec", lambda self: 1)
    assert win.import_changes(pack)
    assert win.ds.get(task, "planEnd") == date(2030, 1, 1)
    results = {"format": "gantt-exchange", "schemaVersion": "1.0",
               "meta": dict(pack["meta"], kind="results"),
               "results": [{"changeId": pack["changes"][0]["changeId"], "result": "applied"}]}
    assert check(results).errors == []
    monkeypatch.setattr(incoming.ResultsDialog, "exec", lambda self: 1)
    assert win.import_results(results)
    assert not win.ds.edit_count and task.source("planEnd") == date(2030, 1, 1)   # теперь это значение источника


def test_incremental(win, root):
    before = len(win.ds.nodes)
    assert win.open_file(str(root / "examples" / "incremental.gantt.json"))
    assert len(win.ds.nodes) == before       # одно новое, одно удалено


def test_clipboard_and_wizard(win, qapp, monkeypatch):
    from gantt.ui import wizard

    qapp.clipboard().setText("№\tНазвание\tСрок\tИсполнитель\n1\tЭтап\t\t\n1.1\tПодэтап\t01.10.2026\tАналитик 1\n")
    monkeypatch.setattr(wizard.TableWizard, "exec", lambda self: 1)
    win.paste_table()
    assert win.ds.counts() == {"group": 0, "project": 1, "task": 1}
    assert win.path is None and win.source_label.startswith("Буфер обмена")


def test_filters_open_over_the_table_and_close(win, qapp):
    from PySide6.QtCore import QEvent, QPointF, Qt
    from PySide6.QtGui import QMouseEvent

    win.show()
    win.settings.animations = False
    win.filters_pop.animations = False
    win.toggle_filters(True)
    assert win.filters_pop.is_open() and win.a_filters.isChecked()
    assert win.filters_pop.parent() is win                     # поверх окна, а не док сбоку
    press = QMouseEvent(QEvent.Type.MouseButtonPress, QPointF(5, 5), QPointF(5, 5), Qt.MouseButton.LeftButton,
                        Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    win.filters_pop.eventFilter(win.timeline.viewport(), press)  # щелчок мимо — закрылась
    assert not win.filters_pop.is_open() and not win.a_filters.isChecked()


def test_gap_collapses_when_column_hidden(win):
    win.show()
    win.settings.animations = False
    t = win.table
    total = sum(win.splitter.sizes())
    win.splitter.setSizes([t.content_width(), total - t.content_width()])
    win._toggle_column("executors", False)
    win._collapse_gap()
    assert win.splitter.sizes()[0] <= t.content_width() + 1      # пустого места между таблицей и шкалой нет


def test_hover_card_explains_plan_and_fact(win):
    from gantt.ui.hovercard import card_lines

    node = next(n for n in win.ds.walk() if win.ds.role(n) == "task" and n.raw.get("factStart"))
    lines = card_lines(win.ds, node, win.settings, win.tokens, win.project_color)
    labels = [ln.label for ln in lines if ln.kind == "row"]
    assert "План" in labels and "Факт" in labels
    assert lines[0].text == node.name
    win.card._set(node)
    assert win.card.height() > 120


def test_search_restores_expansion(win):
    from gantt.ui.tree_model import NodeRole

    before = {i.data(NodeRole).id for i in win._all_indexes() if win.table.isExpanded(i)}
    win.search.setText("ERP")
    win.apply_filter()
    assert win.proxy.accepted
    win.search.setText("")
    win.apply_filter()
    after = {i.data(NodeRole).id for i in win._all_indexes() if win.table.isExpanded(i)}
    assert after == before


def test_ctrl_click_groups_columns_and_toggle_hides_them(win, qapp):
    from PySide6.QtCore import QEvent, QPointF, Qt
    from PySide6.QtGui import QMouseEvent

    from gantt.ui.tree_model import COL

    win.show()
    win.settings.animations = False
    h = win.table.header()

    def ctrl_click(key):
        x = h.sectionViewportPosition(COL[key]) + 10
        ev = QMouseEvent(QEvent.Type.MouseButtonPress, QPointF(x, 30), QPointF(x, 30), Qt.MouseButton.LeftButton,
                         Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ControlModifier)
        h.mousePressEvent(ev)

    for key in ("progress", "due"):
        win.table.setColumnHidden(COL[key], False)
    ctrl_click("due")
    ctrl_click("progress")
    assert h.selection == ["due", "progress"]
    h._check_ctrl()                                    # Ctrl отпущен — группа готова
    assert win.settings.column_groups == [{"cols": ["progress", "due"], "collapsed": False}]
    assert h.visualIndex(COL["due"]) == h.visualIndex(COL["progress"]) + 1   # стали рядом
    win.toggle_group(0)
    assert win.table.isColumnHidden(COL["progress"]) and win.table.isColumnHidden(COL["due"])
    assert "progress" not in win.settings.hidden_columns        # меню колонок не трогаем
    win.toggle_group(0)
    assert not win.table.isColumnHidden(COL["progress"])
    win.ungroup(0)
    assert win.settings.column_groups == []


def test_panels_open_under_their_buttons_one_at_a_time(win):
    win.show()
    for pop in win.pops:
        pop.animations = False
    win.toggle_theme()
    assert win.theme_pop.is_open() and win.theme_btn.isChecked()
    win.toggle_changes(True)
    assert win.changes_pop.is_open() and not win.theme_pop.is_open()   # одна панель за раз
    win.toggle_filters(True)
    assert win.filters_pop.is_open() and not win.changes_pop.is_open()
    # в фильтрах нет общей прокрутки: панель по размеру содержимого, вбок списки не крутятся
    from PySide6.QtWidgets import QListWidget, QScrollArea
    assert not win.filters.findChildren(QScrollArea)
    assert win.filters_pop.width() >= win.filters.sizeHint().width()
    for lst in win.filters.findChildren(QListWidget):
        assert lst.horizontalScrollBarPolicy().name == "ScrollBarAlwaysOff"


def test_level_buttons_expand_to_level(win):
    from gantt.ui.tree_model import NodeRole

    h = win.table.header()
    assert h.levels == 3 and h.level_names[0] == "Раздел"
    win.expand_to_level(1)
    assert not any(win.table.isExpanded(i) for i in win._all_indexes())
    win.expand_to_level(2)
    expanded = [i.data(NodeRole) for i in win._all_indexes() if win.table.isExpanded(i)]
    assert expanded and all(n.depth == 0 for n in expanded) and h.level_active == 2
    win.expand_to_level(3)
    assert all(win.table.isExpanded(i) for i in win._all_indexes() if i.data(NodeRole).children)


def test_hidden_group_columns_leave_no_plus(win):
    from dataclasses import replace

    from gantt.ui.tree_model import COL

    h = win.table.header()
    hidden = [k for k in win.settings.hidden_columns if k not in ("status", "progress")]
    win.apply_settings(replace(win.settings, hidden_columns=hidden,
                               column_groups=[{"cols": ["status", "progress"], "collapsed": True}]), save=False)
    assert [g["index"] for g in h.groups] == [0] and h._bar() == h.BAR
    # все колонки группы выключены в меню — над заголовком ни «+», ни полоски групп
    win.apply_settings(replace(win.settings, hidden_columns=hidden + ["status", "progress"]), save=False)
    assert h.groups == [] and h._bar() == 0
    # включили колонку из свёрнутой группы — группа разворачивается, колонка видна
    win._toggle_column("status", True)
    assert win.settings.column_groups == [{"cols": ["status", "progress"], "collapsed": False}]
    assert not win.table.isColumnHidden(COL["status"]) and win.table.isColumnHidden(COL["progress"])
    assert [g["cols"] for g in h.groups] == [["status"]]


def test_column_list_shows_hides_and_groups(win, qapp):
    from dataclasses import replace

    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtTest import QTest

    from gantt.ui.columns import ColumnList
    from gantt.ui.tree_model import COL

    win.show()
    hidden = [k for k in win.settings.hidden_columns if k not in ("progress", "due")]
    win.apply_settings(replace(win.settings, hidden_columns=hidden, column_groups=[]), save=False)
    closed = []
    cl = ColumnList(win, lambda: closed.append(True))
    cl.resize(cl.sizeHint())
    cl.show()
    rows, rh = cl.rows, cl.rows.row_h()

    def click(key, mods=Qt.KeyboardModifier.NoModifier):
        y = cl.keys().index(key) * rh + rh // 2
        QTest.mouseClick(rows, Qt.MouseButton.LeftButton, mods, QPoint(60, y))

    click("number")                                        # щелчок — показать или скрыть, меню не закрывается
    assert ("number" in win.settings.hidden_columns) != ("number" in hidden) and not closed
    click("progress", Qt.KeyboardModifier.ControlModifier)  # Ctrl+щелчок — отметить для группы
    assert cl.picking and cl.marked == ["progress"]
    click("due")                                           # в режиме отметки — просто щелчок
    click("start")                                         # скрытую колонку в группу не берём
    assert cl.marked == ["progress", "due"]
    cl.btn_group.click()
    assert closed and win.settings.column_groups == [{"cols": ["progress", "due"], "collapsed": False}]
    # «−» у скобки группы сворачивает её, «×» — разгруппировывает
    cl2 = ColumnList(win)
    cl2.resize(cl2.sizeHint())
    cl2.rows.grab()
    fold = next(r for r, kind, g in cl2.rows._buttons if kind == "fold")
    QTest.mouseClick(cl2.rows, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, fold.center().toPoint())
    assert win.settings.column_groups[0]["collapsed"] and win.table.isColumnHidden(COL["due"])
    drop = next(r for r, kind, g in cl2.rows._buttons if kind == "drop")
    QTest.mouseClick(cl2.rows, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier, drop.center().toPoint())
    assert win.settings.column_groups == [] and not win.table.isColumnHidden(COL["due"])


def test_font_size_reaches_widgets_and_panels_fit(make_win, qapp):
    from dataclasses import replace

    from gantt.settings import load_view, save_view
    from gantt.ui.legend import _Line, ITEMS

    save_view(replace(load_view(), font_size=12))
    w = make_win()
    try:
        assert w.table.font().pointSize() == 12 and w.legend.btn.font().pointSize() == 12   # с первого запуска
        w.show()
        w.stack.setCurrentIndex(1)
        for pop in w.pops:
            pop.animations = False
        w.toggle_theme()                               # панель уже открывали — размеры запомнены
        w.toggle_theme()
        w.apply_settings(replace(w.settings, font_size=14), save=False)
        qapp.processEvents()
        assert w.legend.btn.width() >= w.legend.btn.sizeHint().width()        # «Все обозначения…» без многоточия
        w.toggle_theme()                               # панель была скрыта, пока менялся шрифт
        qapp.processEvents()
        for b in w.theme_panel.themes.buttons():
            assert b.height() >= b.sizeHint().height()                       # подписи тем не обрезаны
        assert len(w.legend_panel.findChildren(_Line)) == len(ITEMS) - 1     # обозначения разделены чертой
    finally:
        w.apply_settings(replace(w.settings, font_size=10), save=False)


def test_settings_panel_groups_columns_and_keeps_them(win, qapp):
    from dataclasses import replace

    from PySide6.QtCore import QPoint, Qt
    from PySide6.QtTest import QTest

    win.show()
    for pop in win.pops:
        pop.animations = False
    hidden = [k for k in win.settings.hidden_columns if k not in ("progress", "due")]
    win.apply_settings(replace(win.settings, hidden_columns=hidden, column_groups=[]), save=False)
    win.open_settings()
    assert win.settings_pop.is_open()
    cl = win.settings_panel.columns
    rh = cl.rows.row_h()

    def click(key, mods=Qt.KeyboardModifier.NoModifier):
        QTest.mouseClick(cl.rows, Qt.MouseButton.LeftButton, mods, QPoint(60, cl.keys().index(key) * rh + rh // 2))

    click("progress", Qt.KeyboardModifier.ControlModifier)
    click("due")
    cl.btn_group.click()
    assert win.settings.column_groups == [{"cols": ["progress", "due"], "collapsed": False}]
    assert win.settings_pop.is_open()                  # панель остаётся открытой
    click("number")
    number_hidden = "number" in win.settings.hidden_columns
    win.settings_panel.zebra.toggle()                  # другая настройка не откатывает колонки и группы
    assert win.settings.column_groups == [{"cols": ["progress", "due"], "collapsed": False}]
    assert ("number" in win.settings.hidden_columns) == number_hidden
