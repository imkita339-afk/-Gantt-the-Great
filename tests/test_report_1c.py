"""Excel-отчёт 1С ProjectsToGanttExport: синтетический файл всегда, реальный — если лежит в local/."""
from pathlib import Path

import openpyxl
import pytest

from gantt.exchange import OpenError
from gantt.formats import NeedsWizard, open_path
from gantt.model import Dataset

LOCAL = Path(__file__).resolve().parent.parent / "local"
REAL = next(iter(sorted(LOCAL.glob("ProjectsToGanttExport*.xlsx"))), None)  # пример пользователя — только локально
EXPECTED = LOCAL / "expected.json"  # ожидаемые числа по нему — тоже только локально


def make_report(path: Path, with_task_dates: bool = True) -> Path:
    """Мини-отчёт той же структуры: шапка в строке 4, группировка строк задаёт уровни."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"], ws["C2"] = "Параметры:", "Значение: —"
    head = {"A": "Проект", "E": "Дата начала план", "F": "Дата окончания план", "G": "Дата начала",
            "H": "Дата окончания", "J": "Задание", "L": "Текущий исполнитель"}
    if with_task_dates:
        head.update({"M": "Дата", "N": "Крайняя дата выполнения", "O": "Дата выполнения"})
    for col, text in head.items():
        ws[f"{col}4"] = text
    rows = [
        (0, {"A": "1.Первый раздел"}),
        (1, {"A": "ERP. Проект", "E": "01.10.2026", "F": "30.11.2026"}),
        (2, {"A": "ERP. Проект", "J": "Инцидент № 101 от 01.09.26 (Первая задача)", "L": "Аналитик 1",
             "M": "01.09.2026 10:00:00", "O": "10.09.2026 9:08:45"}),
        (2, {"A": "ERP. Проект", "J": "Запрос на изменение № 102 от 02.09.26 (Вторая задача...)",
             "L": "Разработчик 2", "M": "02.09.2026 11:00:00", "N": "20.09.2026 0:00:00"}),
        (2, {"A": "ERP. Проект", "J": "Инцидент № 103 от 03.09.26 (Третья задача)", "L": "Аналитик 1",
             "M": "03.09.2026 12:00:00"}),
        (1, {"A": "Пустой проект"}),
        (2, {"A": "Пустой проект"}),
        (0, {"A": "Архив"}),
        (1, {"A": "Старый проект", "E": "01.01.2026", "F": "31.03.2026", "G": "01.01.2026", "H": "15.03.2026"}),
        (2, {"A": "Старый проект", "J": "Консультация № 90 от 05.01.26 (Старое)", "L": "Тестировщик 3",
             "M": "05.01.2026 9:00:00", "O": "06.01.2026 9:00:00"}),
    ]
    for i, (lvl, cells) in enumerate(rows, start=5):
        for col, value in cells.items():
            if col in ("M", "N", "O") and not with_task_dates:
                continue
            ws[f"{col}{i}"] = value
        ws.row_dimensions[i].outline_level = lvl
    wb.save(path)
    return path


def test_synthetic_report(tmp_path):
    res = open_path(make_report(tmp_path / "ProjectsToGanttExport2026-01-15_09-30.xlsx"))
    doc = res.document
    ds = Dataset(doc)
    assert ds.counts() == {"group": 2, "project": 3, "task": 4}
    sec = ds.roots[0]
    assert (sec.number, sec.name) == ("1", "Первый раздел")
    proj = sec.children[0]
    assert ds.progress(proj) == (1, 3)
    assert ds.progress_label(proj) == "33% · 1/3"
    t1, t2, t3 = proj.children
    assert (t1.number, t1.name) == ("101", "Первая задача")
    assert ds.status_label(t1) == "Выполнено"
    assert ds.task_span(t3)[2] is True           # нет ни срока, ни выполнения — открытый конец
    assert doc["meta"]["sourceName"] == "ProjectsToGanttExport"
    assert any("без заданий и без дат: 1" in w for w in res.warnings)
    assert any("обрезанных" in w for w in res.warnings)


def test_overdue_by_deadline(tmp_path):
    from datetime import date
    doc = open_path(make_report(tmp_path / "r.xlsx")).document
    ds = Dataset(doc, today=date(2026, 9, 24))
    t2 = ds.roots[0].children[0].children[1]
    assert ds.is_overdue(t2) and ds.status_label(t2) == "Просрочено"


def test_old_report_without_task_dates(tmp_path):
    res = open_path(make_report(tmp_path / "old-report.xlsx", with_task_dates=False))
    assert any("нет колонок" in w for w in res.warnings)


def test_not_a_report(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Что-то другое"
    path = tmp_path / "other.xlsx"
    wb.save(path)
    with pytest.raises(NeedsWizard):      # незнакомая таблица — открывается мастером
        open_path(path)


def test_ids_are_stable(tmp_path):
    a = open_path(make_report(tmp_path / "a.xlsx")).document
    b = open_path(make_report(tmp_path / "b.xlsx")).document
    assert [n["id"] for n in a["nodes"]] == [n["id"] for n in b["nodes"]]


@pytest.mark.skipif(REAL is None or not EXPECTED.exists(), reason="примера пользователя нет (он только в local/)")
def test_real_report_numbers():
    import json
    expected = json.loads(EXPECTED.read_text(encoding="utf-8"))
    ds = Dataset(open_path(REAL).document)
    assert ds.counts() == expected["counts"]
    with_something = [n for n in ds.walk() if ds.role(n) == "project"
                      and (n.children or any(n.source(f) for f in ("planStart", "planEnd", "factStart", "factEnd")))]
    assert len(with_something) == expected["projects_with_data"]
