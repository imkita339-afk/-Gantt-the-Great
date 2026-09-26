"""«Сохранить как» с выбором формата и честным «что не сохранится», лист изменений (Excel, PDF)."""
from __future__ import annotations

from datetime import datetime
from html import escape

from PySide6.QtCore import QMarginsF
from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument
from PySide6.QtWidgets import (QButtonGroup, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFormLayout, QFrame,
                               QGroupBox, QLabel, QLineEdit, QRadioButton, QStackedWidget, QVBoxLayout,
                               QWidget)

from ..exchange import FIELD_TITLES
from ..formats import WRITERS
from ..model import Dataset
from .panels import value_text


class SaveAsDialog(QDialog):
    """Формат выбирается здесь; для каждого сказано, что в нём не сохранится."""

    def __init__(self, qs, title: str, parent=None):
        super().__init__(parent)
        self.qs = qs
        self.setWindowTitle("Сохранить как")
        self.setMinimumWidth(620)
        lay = QVBoxLayout(self)
        lay.addWidget(QLabel("<b>В каком виде сохранить?</b>"))
        self.group = QButtonGroup(self)
        self.radios: dict[str, QRadioButton] = {}
        last = qs.value("saveAs/format", "json")
        for key, name, ext, desc, loss in WRITERS:
            card = QFrame()
            card.setObjectName("card")
            cl = QVBoxLayout(card)
            cl.setContentsMargins(6, 4, 6, 4)
            cl.setSpacing(1)
            rb = QRadioButton(f"{name}   ({ext})")
            f = rb.font()
            f.setBold(True)
            rb.setFont(f)
            self.group.addButton(rb)
            self.radios[key] = rb
            cl.addWidget(rb)
            d = QLabel(desc)
            d.setWordWrap(True)
            d.setContentsMargins(22, 0, 0, 0)
            cl.addWidget(d)
            lo = QLabel(loss or "Сохраняется всё.")
            lo.setWordWrap(True)
            lo.setContentsMargins(22, 0, 0, 0)
            lo.setStyleSheet("color: #B26A00;" if loss else "color: #2E8B57;")
            cl.addWidget(lo)
            lay.addWidget(card)
            rb.toggled.connect(lambda on, k=key: on and self._format_changed(k))

        self.options = QStackedWidget()
        self.opt_pages = {}
        empty = QWidget()
        self.opt_pages["none"] = self.options.addWidget(empty)

        xl = QGroupBox("Excel")
        xll = QVBoxLayout(xl)
        self.chart = QCheckBox("Добавить лист «Диаграмма» — цветные клетки по неделям")
        self.chart.setChecked(qs.value("saveAs/chart", "true") == "true")
        xll.addWidget(self.chart)
        self.opt_pages["xlsx"] = self.options.addWidget(xl)

        pic = QGroupBox("Картинка")
        form = QFormLayout(pic)
        self.scope = QComboBox()
        self.scope.addItem("Как на экране — с фильтрами и раскрытием", "visible")
        self.scope.addItem("Всё — раскрыто полностью, без фильтров", "all")
        self.scope.setCurrentIndex(max(0, self.scope.findData(qs.value("saveAs/scope", "visible"))))
        form.addRow("Строки:", self.scope)
        self.period = QComboBox()
        self.period.addItem("Весь период этих строк", "data")
        self.period.addItem("Только видимый на экране участок шкалы", "screen")
        self.period.setCurrentIndex(max(0, self.period.findData(qs.value("saveAs/period", "data"))))
        form.addRow("Период:", self.period)
        self.page = QComboBox()
        self.page.addItem("A4, альбомный", "A4")
        self.page.addItem("A3, альбомный", "A3")
        self.page.setCurrentIndex(max(0, self.page.findData(qs.value("saveAs/page", "A4"))))
        self.page_label = QLabel("Лист:")
        form.addRow(self.page_label, self.page)
        self.title = QLineEdit(title)
        form.addRow("Заголовок:", self.title)
        self.opt_pages["picture"] = self.options.addWidget(pic)
        lay.addWidget(self.options)

        bb = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        bb.button(QDialogButtonBox.StandardButton.Ok).setText("Выбрать файл и сохранить…")
        bb.button(QDialogButtonBox.StandardButton.Ok).setObjectName("primary")
        bb.button(QDialogButtonBox.StandardButton.Cancel).setText("Отмена")
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)
        (self.radios.get(last) or self.radios["json"]).setChecked(True)
        self._format_changed(self.key())

    def key(self) -> str:
        return next(k for k, rb in self.radios.items() if rb.isChecked())

    def _format_changed(self, key: str) -> None:
        page = "xlsx" if key == "xlsx" else "picture" if key in ("pdf", "png") else "none"
        self.options.setCurrentIndex(self.opt_pages[page])
        self.page.setVisible(key == "pdf")
        self.page_label.setVisible(key == "pdf")

    def values(self) -> dict:
        v = {"format": self.key(), "chart": self.chart.isChecked(), "scope": self.scope.currentData(),
             "period": self.period.currentData(), "page": self.page.currentData(), "title": self.title.text()}
        self.qs.setValue("saveAs/format", v["format"])
        self.qs.setValue("saveAs/chart", "true" if v["chart"] else "false")
        for k in ("scope", "period", "page"):
            self.qs.setValue(f"saveAs/{k}", v[k])
        return v


# ---------- лист изменений ----------
def change_rows(ds: Dataset) -> list[list[str]]:
    """Где, что, поле, было (в источнике), стало (правка)."""
    rows = []
    for node in ds.walk():
        for fld in sorted(ds.edited_fields(node)):
            path, p = [], node.parent
            while p is not None:
                path.insert(0, p.name)
                p = p.parent
            rows.append([" / ".join(path), node.name, ds.number_label(node), FIELD_TITLES.get(fld, fld),
                         value_text(ds, fld, node.source(fld)), value_text(ds, fld, ds.overrides[(node.id, fld)])])
    return rows


HEAD = ["Раздел / проект", "Что", "№", "Поле", "Было (в источнике)", "Стало"]


def write_change_list_xlsx(ds: Dataset, path: str, author: str = "") -> int:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = change_rows(ds)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист изменений"
    ws.append([f"Лист изменений — {ds.title}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Сформирован {datetime.now():%d.%m.%Y %H:%M}" + (f" · {author}" if author else "") +
               f" · правок: {len(rows)}"])
    ws.append([])
    ws.append(HEAD)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EFEDE7")
    for r in rows:
        ws.append(r)
        ws.cell(ws.max_row, 6).font = Font(bold=True)
    for col, w in zip("ABCDEF", (42, 48, 16, 18, 20, 20)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{max(4, ws.max_row)}"
    wb.save(path)
    return len(rows)


def write_change_list_pdf(ds: Dataset, path: str, author: str = "") -> int:
    rows = change_rows(ds)
    cells = "".join("<tr>" + "".join(f"<td>{escape(str(x))}</td>" for x in r[:-1]) +
                    f"<td><b>{escape(str(r[-1]))}</b></td></tr>" for r in rows)
    html = f"""<html><body style="font-family:'Segoe UI'; font-size:9pt">
    <h2 style="margin:0">Лист изменений — {escape(ds.title)}</h2>
    <p style="color:#666">Сформирован {datetime.now():%d.%m.%Y %H:%M}{escape(' · ' + author) if author else ''}
    · правок: {len(rows)}</p>
    <table border="1" cellspacing="0" cellpadding="4" width="100%" style="border-collapse:collapse;border-color:#ccc">
    <tr style="background:#EFEDE7">{''.join(f'<th align="left">{h}</th>' for h in HEAD)}</tr>{cells}</table>
    <p style="color:#666; margin-top:16px">Подпись: ______________________</p>
    </body></html>"""
    writer = QPdfWriter(path)
    writer.setTitle(f"Лист изменений — {ds.title}")
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.PageSizeId.A4), QPageLayout.Orientation.Landscape,
                                     QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter))
    doc = QTextDocument()
    doc.setHtml(html)
    doc.print_(writer)
    return len(rows)
