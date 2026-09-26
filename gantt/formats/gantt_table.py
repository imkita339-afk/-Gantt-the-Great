"""«Гант-таблица»: сохранение в Excel и CSV так, чтобы программа открывала это обратно."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, timedelta
from pathlib import Path

HEAD = ["Уровень", "№", "Наименование", "Исполнители", "Инициатор", "План начала", "План окончания",
        "Факт начала", "Факт окончания", "Статус", "%", "Приоритет", "Срочность", "Влияние", "Категория", "Цвет"]
TAIL = ["ID", "Родитель"]
SERVICE_SHEET = "_gantt"
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]


def _tree(doc: dict) -> list[tuple[dict, int]]:
    """Узлы в порядке дерева с глубиной."""
    nodes = doc.get("nodes", [])
    children: dict = {}
    for n in nodes:
        children.setdefault(n.get("parentId"), []).append(n)
    for lst in children.values():
        lst.sort(key=lambda n: n.get("order", 0))
    out: list[tuple[dict, int]] = []

    def walk(pid, depth):
        for n in children.get(pid, []):
            out.append((n, depth))
            walk(n["id"], depth + 1)
    walk(None, 0)
    return out


def _names(doc: dict):
    d = doc.get("dictionaries", {})
    by = lambda key: {x["id"]: x for x in d.get(key, [])}  # noqa: E731
    return by("levels"), by("statuses"), by("people"), by("priorities"), by("urgencies"), by("impacts"), \
        by("categories"), {f["id"]: f for f in d.get("customFields", [])}


def rows_for(doc: dict) -> tuple[list[str], list[list], list[int]]:
    """Шапка, строки и глубины для листа «Данные» / CSV."""
    levels, statuses, people, pri, urg, imp, cat, fields = _names(doc)
    custom = list(fields.values())
    head = HEAD + [f["name"] for f in custom] + TAIL
    rows, depths = [], []
    for n, depth in _tree(doc):
        def nm(dic, key):
            v = n.get(key)
            return dic.get(v, {}).get("name", v) if v else None

        cvals = []
        for f in custom:
            v = (n.get("custom") or {}).get(f["id"])
            if f.get("type") == "enum":
                v = next((x["name"] for x in f.get("values", []) if x["id"] == v), v)
            cvals.append(v)
        rows.append([levels.get(n["level"], {}).get("name", n["level"]), n.get("number"), n["name"],
                     ", ".join(people.get(p, {}).get("name", p) for p in n.get("executors") or []) or None,
                     nm(people, "initiator"),
                     *[date.fromisoformat(n[f]) if n.get(f) else None
                       for f in ("planStart", "planEnd", "factStart", "factEnd")],
                     nm(statuses, "status"), n.get("percent"), nm(pri, "priority"), nm(urg, "urgency"),
                     nm(imp, "impact"), nm(cat, "category"), n.get("color"), *cvals, n["id"], n.get("parentId")])
        depths.append(depth)
    return head, rows, depths


def write_xlsx(doc: dict, path: str | Path, colors: dict | None = None, chart: bool = True) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    colors = colors or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"
    head, rows, depths = rows_for(doc)
    ws.append(head)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="EFEDE7")
    levels = sorted(doc["dictionaries"]["levels"], key=lambda x: x["order"])
    leaf = levels[-1]["name"] if levels else ""
    for row, depth in zip(rows, depths):
        ws.append(row)
        r = ws.max_row
        ws.row_dimensions[r].outline_level = min(depth, 7)
        ws.cell(r, 3).alignment = Alignment(indent=depth)
        for col in (6, 7, 8, 9):
            ws.cell(r, col).number_format = "DD.MM.YYYY"
        nid = row[-2]
        if row[0] != leaf:
            for col in range(1, len(head) + 1):
                ws.cell(r, col).font = bold
            color = colors.get(nid)
            fill = color.lstrip("#") if color else "F0EEE8"
            ws.cell(r, 3).fill = PatternFill("solid", fgColor=fill)
    ws.sheet_properties.outlinePr.summaryBelow = False
    widths = {1: 12, 2: 9, 3: 60, 4: 28, 5: 20, 6: 12, 7: 14, 8: 12, 9: 14, 10: 14, 11: 6}
    for i in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 14)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    ref = wb.create_sheet("Справочники")
    ref.append(["Справочник", "Значение", "Смысл / порядок", "Цвет"])
    for c in ref[1]:
        c.font = bold
    d = doc["dictionaries"]
    for title, key, extra in (("Уровни", "levels", "order"), ("Статусы", "statuses", "kind"),
                              ("Приоритет", "priorities", "order"), ("Срочность", "urgencies", "order"),
                              ("Влияние", "impacts", "order"), ("Категории", "categories", None),
                              ("Люди", "people", "fullName")):
        for x in d.get(key, []):
            ref.append([title, x.get("name"), x.get(extra) if extra else None, x.get("color")])
    ref.column_dimensions["A"].width = 16
    ref.column_dimensions["B"].width = 32
    ref.column_dimensions["C"].width = 30

    if chart:
        _chart_sheet(wb, doc, colors)

    svc = wb.create_sheet(SERVICE_SHEET)
    # сохранённая таблица — свои данные со своими правками, а не тот источник, из которого она сделана:
    # иначе её открытие «приняло» бы правки исходного набора и стёрло их из очереди
    meta = dict(doc["meta"], source="table:Гант-таблица", sourceName=Path(path).name)
    meta.pop("changedSince", None)
    payload = json.dumps({"meta": meta, "dictionaries": doc["dictionaries"]}, ensure_ascii=False)
    svc["A1"] = "gantt-exchange"
    for i in range(0, len(payload), 30000):
        svc.cell(2 + i // 30000, 1, payload[i:i + 30000])
    svc.sheet_state = "hidden"
    wb.save(path)


def _chart_sheet(wb, doc: dict, colors: dict) -> None:
    """Понедельная шкала с цветными ячейками — как привычный Excel-Гант, для печати и руководителя."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    tree = _tree(doc)
    dates = [date.fromisoformat(n[f]) for n, _d in tree for f in ("planStart", "planEnd", "factStart", "factEnd")
             if n.get(f)]
    if not dates:
        return
    start = min(dates)
    start -= timedelta(days=start.weekday())
    weeks = min(160, (max(dates) - start).days // 7 + 2)
    ws = wb.create_sheet("Диаграмма")
    ws.cell(1, 1, "Наименование").font = Font(bold=True)
    for w in range(weeks):
        d = start + timedelta(weeks=w)
        c = ws.cell(1, 2 + w, f"{d.day:02d} {MONTHS[d.month - 1]}")
        c.alignment = Alignment(text_rotation=90, horizontal="center")
        c.font = Font(size=8)
        ws.column_dimensions[get_column_letter(2 + w)].width = 2.6
    ws.column_dimensions["A"].width = 55
    ws.row_dimensions[1].height = 48
    kinds = {x["id"]: x.get("kind") for x in doc["dictionaries"].get("statuses", [])}
    kind_color = {"done": "BFE3CC", "inProgress": "C3D3F3", "overdue": "F7C7A8", "notStarted": "DDE0E5",
                  "onHold": "EDE0B8", "release": "D9CCF2", "cancelled": "E3E3E3"}
    today = date.today()
    for r, (n, depth) in enumerate(tree, start=2):
        c = ws.cell(r, 1, n["name"])
        c.alignment = Alignment(indent=depth)
        ws.row_dimensions[r].outline_level = min(depth, 7)
        s = n.get("planStart") or n.get("factStart")
        e = n.get("planEnd") or n.get("factEnd")
        if not s:
            continue
        s = date.fromisoformat(s)
        e = date.fromisoformat(e) if e else max(s, today)
        kind = kinds.get(n.get("status") or "")
        if kind not in ("done", "cancelled") and n.get("planEnd") and e < today:
            kind = "overdue"
        fill = (colors.get(n["id"]) or "").lstrip("#") or kind_color.get(kind or "", "C3D3F3")
        for w in range(weeks):
            ws_start = start + timedelta(weeks=w)
            if ws_start <= e and ws_start + timedelta(days=6) >= s:
                ws.cell(r, 2 + w).fill = PatternFill("solid", fgColor=fill)
    tw = (today - start).days // 7
    if 0 <= tw < weeks:
        ws.cell(1, 2 + tw).font = Font(size=8, bold=True, color="C2185B")
    ws.freeze_panes = "B2"


def write_csv(doc: dict, path: str | Path) -> None:
    head, rows, _depths = rows_for(doc)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(head)
    for row in rows:
        w.writerow([v.strftime("%d.%m.%Y") if isinstance(v, date) else ("" if v is None else v) for v in row])
    Path(path).write_text(buf.getvalue(), encoding="utf-8-sig")


def is_gantt_table(headers: list) -> bool:
    h = [str(x).strip() for x in headers if x not in (None, "")]
    return h[:3] == HEAD[:3] and h[-2:] == TAIL


def read_xlsx(path: str | Path) -> dict | None:
    """Excel, сохранённый программой: восстановить документ без мастера. None — это не «Гант-таблица»."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    if SERVICE_SHEET not in wb.sheetnames or "Данные" not in wb.sheetnames:
        return None
    svc = wb[SERVICE_SHEET]
    if svc["A1"].value != "gantt-exchange":
        return None
    payload = "".join(str(svc.cell(r, 1).value or "") for r in range(2, svc.max_row + 1))
    head = json.loads(payload)
    rows = [list(r) for r in wb["Данные"].iter_rows(values_only=True)]
    return _rebuild(head["meta"], head["dictionaries"], rows)


def read_rows(rows: list[list], name: str) -> dict:
    """CSV «Гант-таблицы»: справочники собираются из значений."""
    from .table import guess_status_kind, scale_order, short_name  # noqa: F401

    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    level_names: list[str] = []
    statuses, scales, people = {}, {"Приоритет": {}, "Срочность": {}, "Влияние": {}, "Категория": {}}, {}
    for r in rows[1:]:
        get = lambda h: r[idx[h]] if h in idx and idx[h] < len(r) and r[idx[h]] not in ("", None) else None  # noqa
        lv = get("Уровень")
        if lv and lv not in level_names:
            level_names.append(lv)
        st = get("Статус")
        if st:
            statuses.setdefault(st, guess_status_kind(st))
        for k in scales:
            v = get(k)
            if v:
                scales[k].setdefault(v, len(scales[k]) + 1)
        for p in (get("Исполнители") or "").split(","):
            if p.strip():
                people.setdefault(p.strip(), None)
        if get("Инициатор"):
            people.setdefault(get("Инициатор"), None)
    fixed = set(HEAD) | set(TAIL)
    custom_names = [h for h in header if h and h not in fixed]
    d = {"levels": [{"id": f"L{i + 1}", "name": n, "order": i + 1} for i, n in enumerate(level_names)],
         "statuses": [{"id": f"st{i + 1}", "name": n, "kind": k, "order": i + 1}
                      for i, (n, k) in enumerate(statuses.items())],
         "people": [{"id": f"p{i + 1}", "name": n} for i, n in enumerate(people)]}
    for title, key in (("Приоритет", "priorities"), ("Срочность", "urgencies"), ("Влияние", "impacts")):
        if scales[title]:
            d[key] = [{"id": f"{key[:3]}{i + 1}", "name": n, "order": scale_order(n, pos)}
                      for i, (n, pos) in enumerate(scales[title].items())]
    if scales["Категория"]:
        d["categories"] = [{"id": f"cat{i + 1}", "name": n} for i, n in enumerate(scales["Категория"])]
    if custom_names:
        d["customFields"] = [{"id": f"c{i + 1}", "name": n, "type": "string"} for i, n in enumerate(custom_names)]
    from datetime import datetime
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    meta = {"source": "table:Гант-таблица", "sourceName": name, "title": name, "exportedAt": stamp,
            "kind": "full", "cursor": stamp}
    return _rebuild(meta, d, rows)


def _rebuild(meta: dict, d: dict, rows: list[list]) -> dict:
    from .table import parse_date, parse_percent

    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    rev = lambda key: {x["name"]: x["id"] for x in d.get(key, [])}  # noqa: E731
    levels, statuses, people = rev("levels"), rev("statuses"), rev("people")
    scales = {"priority": rev("priorities"), "urgency": rev("urgencies"), "impact": rev("impacts"),
              "category": rev("categories")}
    fields = {f["name"]: f for f in d.get("customFields", [])}
    nodes = []
    for order, r in enumerate(rows[1:]):
        def get(h):
            i = idx.get(h)
            return r[i] if i is not None and i < len(r) and r[i] not in ("", None) else None

        if not get("ID") or not get("Наименование"):
            continue
        n = {"id": str(get("ID")), "parentId": str(get("Родитель")) if get("Родитель") else None,
             "level": levels.get(get("Уровень"), next(iter(levels.values()), "L1")), "order": order,
             "name": str(get("Наименование"))}
        if get("№") is not None:
            n["number"] = str(get("№"))
        for h, f in (("План начала", "planStart"), ("План окончания", "planEnd"), ("Факт начала", "factStart"),
                     ("Факт окончания", "factEnd")):
            v = parse_date(get(h))
            if v:
                n[f] = v
        ex = [people[p.strip()] for p in str(get("Исполнители") or "").split(",") if p.strip() in people]
        if ex:
            n["executors"] = ex
        if get("Инициатор") in people:
            n["initiator"] = people[get("Инициатор")]
        if get("Статус") in statuses:
            n["status"] = statuses[get("Статус")]
        pct = parse_percent(get("%"))
        if pct is not None:
            n["percent"] = pct
        for h, f in (("Приоритет", "priority"), ("Срочность", "urgency"), ("Влияние", "impact"),
                     ("Категория", "category")):
            if get(h) in scales[f]:
                n[f] = scales[f][get(h)]
        if get("Цвет"):
            n["color"] = str(get("Цвет")).upper()
        custom = {}
        for name, f in fields.items():
            v = get(name)
            if v is None:
                continue
            if f.get("type") == "enum":
                v = next((x["id"] for x in f.get("values", []) if x["name"] == v), None)
            elif f.get("type") == "number":
                try:
                    v = float(v) if not isinstance(v, (int, float)) else v
                except ValueError:
                    v = None
            elif not isinstance(v, str):
                v = str(v) if f.get("type") == "string" else v
            if v is not None:
                custom[f["id"]] = v
        if custom:
            n["custom"] = custom
        nodes.append(n)
    ids = {n["id"] for n in nodes}
    for n in nodes:
        if n["parentId"] not in ids:
            n["parentId"] = None
    return {"format": "gantt-exchange", "schemaVersion": "1.0", "meta": dict(meta, kind="full"),
            "dictionaries": d, "nodes": nodes}
