"""Встроенный шаблон: Excel-отчёт 1С «ProjectsToGanttExport».

Иерархия — группировка строк Excel: 0 — раздел, 1 — проект, 2 — строка задания. Колонки ищутся по заголовкам:
в разных версиях отчёта их набор отличается, недостающие — не ошибка. Правила — docs/import-export.md.
"""
from __future__ import annotations

import re
import uuid
from datetime import date, datetime
from pathlib import Path

import openpyxl

from ..exchange import OpenError, check
from . import ReadResult

NS = uuid.uuid5(uuid.NAMESPACE_URL, "urn:gantt-exchange:excel:ProjectsToGanttExport")
SOURCE = "excel:ProjectsToGanttExport"
HEADERS = {
    "Проект": "project", "Дата начала план": "planStart", "Дата окончания план": "planEnd",
    "Дата начала": "factStart", "Дата окончания": "factEnd", "Задание": "task",
    "Текущий исполнитель": "executor", "Дата": "created",
    "Крайняя дата выполнения": "deadline", "Дата выполнения": "completed",
}
TASK_RE = re.compile(r"^(?P<kind>.+?) № (?P<num>\S+) от (?P<d>\d{2}\.\d{2}\.\d{2}) \((?P<subject>.*)\)$", re.S)
SECTION_RE = re.compile(r"^(\d+)\.\s*(.+)$")
KIND_IDS = {"Инцидент": "incident", "Запрос на изменение": "changeRequest", "Консультация": "consultation",
            "Задание": "assignment", "Запрос на обслуживание": "serviceRequest"}
EMPTY = {"", "—", "-", "–"}


def gid(*parts: str) -> str:
    return str(uuid.uuid5(NS, "|".join(parts)))


def parse_date(v) -> str | None:
    """Дата из ячейки: дата Excel или текст «дд.мм.гггг», «дд.мм.гг», «ГГГГ-ММ-ДД», с временем или без."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if s in EMPTY:
        return None
    s = s.split(" ")[0]
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def short_name(full: str) -> str:
    p = full.split()
    return f"{p[0]} {p[1][0]}.{p[2][0]}." if len(p) == 3 else full


def find_header(ws) -> tuple[int, dict[str, int]] | None:
    for r in range(1, min(ws.max_row, 20) + 1):
        cols = {}
        for c in ws[r]:
            t = str(c.value).strip() if c.value is not None else ""
            if t in HEADERS:
                cols[HEADERS[t]] = c.column
        if "project" in cols and "task" in cols:
            return r, cols
    return None


def read(path: Path) -> ReadResult:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # повреждённый файл, не xlsx, открыт с паролем
        raise OpenError("Не удалось прочитать Excel-файл.", [str(e)])
    ws = wb.worksheets[0]
    found = find_header(ws)
    if not found:
        raise OpenError("Эта таблица не похожа на отчёт 1С ProjectsToGanttExport.",
                        ["Не найдены колонки «Проект» и «Задание».",
                         "Другие таблицы Excel откроются через мастер сопоставления колонок — в следующем шаге."])
    hdr_row, cols = found
    doc, warnings = convert(ws, hdr_row, cols, path)
    res = check(doc)
    if res.errors:
        raise OpenError("Отчёт прочитан, но результат не прошёл проверку.", res.errors)
    return ReadResult(doc, warnings + res.warnings, "Excel · отчёт 1С ProjectsToGanttExport")


def convert(ws, hdr_row: int, cols: dict[str, int], path: Path) -> tuple[dict, list[str]]:
    def cell(r, key):
        c = cols.get(key)
        v = ws.cell(r, c).value if c else None
        return v.strip() if isinstance(v, str) else v

    people: dict[str, dict] = {}
    kinds: dict[str, str] = {}
    nodes: list[dict] = []
    section_id = project_id = None
    order = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        lvl = ws.row_dimensions[r].outline_level or 0
        name = cell(r, "project")
        order += 1
        if lvl == 0 and name:
            m = SECTION_RE.match(str(name))
            number, title = (m[1], m[2]) if m else (None, str(name))
            section_id = gid("section", title)
            project_id = None
            nodes.append({"id": section_id, "parentId": None, "level": "section", "order": order,
                          "number": number, "name": title})
        elif lvl == 1 and name:
            project_id = gid("project", str(name))
            nodes.append({"id": project_id, "parentId": section_id, "level": "project", "order": order,
                          "name": str(name), "planStart": parse_date(cell(r, "planStart")),
                          "planEnd": parse_date(cell(r, "planEnd")), "factStart": parse_date(cell(r, "factStart")),
                          "factEnd": parse_date(cell(r, "factEnd"))})
        elif lvl >= 2 and cell(r, "task") and project_id:
            text = str(cell(r, "task"))
            m = TASK_RE.match(text)
            kind, num, subject = (m["kind"], m["num"], m["subject"]) if m else (None, None, text)
            executors = []
            ex = cell(r, "executor")
            if ex:
                pid = gid("person", str(ex))
                people.setdefault(pid, {"id": pid, "name": short_name(str(ex)), "fullName": str(ex)})
                executors.append(pid)
            created, deadline, completed = (parse_date(cell(r, k)) for k in ("created", "deadline", "completed"))
            node = {"id": gid("task", num or text), "parentId": project_id, "level": "task", "order": order,
                    "number": num, "name": subject, "executors": executors,
                    "planStart": created, "planEnd": deadline, "factStart": created, "factEnd": completed,
                    "status": "done" if completed else "inProgress"}
            if kind:
                kinds.setdefault(kind, KIND_IDS.get(kind) or gid("kind", kind))
                node["custom"] = {"taskKind": kinds[kind]}
            nodes.append(node)

    m = re.search(r"(\d{4}-\d{2}-\d{2})_(\d{2})-(\d{2})", path.name)
    stamp = datetime.fromisoformat(f"{m[1]}T{m[2]}:{m[3]}:00") if m else datetime.fromtimestamp(path.stat().st_mtime)
    exported = stamp.astimezone().isoformat(timespec="seconds")
    doc = {
        "format": "gantt-exchange",
        "schemaVersion": "1.0",
        "meta": {
            "source": SOURCE,
            # набор данных — сам отчёт, а не конкретный файл: правки переживают новые выгрузки
            "sourceName": "ProjectsToGanttExport",
            "title": f"Проекты и задания · отчёт от {stamp:%d.%m.%Y %H:%M}",
            "exportedAt": exported,
            "kind": "full",
            "cursor": exported,
        },
        "dictionaries": {
            "levels": [{"id": "section", "name": "Раздел", "order": 1},
                       {"id": "project", "name": "Проект", "order": 2},
                       {"id": "task", "name": "Задание", "order": 3}],
            "statuses": [{"id": "inProgress", "name": "В работе", "kind": "inProgress", "order": 1},
                         {"id": "done", "name": "Выполнено", "kind": "done", "order": 2}],
            "people": sorted(people.values(), key=lambda p: p["name"]),
            "customFields": [{"id": "taskKind", "name": "Вид задания", "type": "enum",
                              "values": [{"id": v, "name": k} for k, v in kinds.items()] or
                                        [{"id": "task", "name": "Задание"}]}],
        },
        "nodes": nodes,
    }
    warnings = []
    by_parent: dict[str, int] = {}
    for n in nodes:
        if n["parentId"]:
            by_parent[n["parentId"]] = by_parent.get(n["parentId"], 0) + 1
    empty = [n for n in nodes if n["level"] == "project" and not by_parent.get(n["id"])
             and not any(n.get(k) for k in ("planStart", "planEnd", "factStart", "factEnd"))]
    if empty:
        warnings.append(f"Проектов без заданий и без дат: {len(empty)} — показаны без полосы")
    missing = [t for t, k in (("«Дата»", "created"), ("«Крайняя дата выполнения»", "deadline"),
                              ("«Дата выполнения»", "completed")) if k not in cols]
    if missing:
        warnings.append(f"В отчёте нет колонок {', '.join(missing)} — у заданий не будет части дат")
    return doc, warnings
