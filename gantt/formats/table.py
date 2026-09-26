"""Любая таблица (Excel, CSV, буфер обмена) → файл обмена по сопоставлению колонок (мастер или шаблон)."""
from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

NS = uuid.uuid5(uuid.NAMESPACE_URL, "urn:gantt-exchange:table")

FIELDS = [  # (поле, название для человека)
    ("name", "Наименование"), ("number", "№ / код"), ("key", "Ключ строки"), ("parent", "Родитель"),
    ("level", "Уровень"), ("executors", "Исполнители"), ("initiator", "Инициатор"),
    ("planStart", "План: начало"), ("planEnd", "План: окончание"), ("factStart", "Факт: начало"),
    ("factEnd", "Факт: окончание"), ("percent", "Процент"), ("status", "Статус"), ("priority", "Приоритет"),
    ("urgency", "Срочность"), ("impact", "Влияние"), ("category", "Категория"), ("color", "Цвет"),
]
FIELD_TITLES = dict(FIELDS)
SYNONYMS = {
    "name": ["наименование", "название", "тема", "задача", "задание", "наименование задачи", "name", "title",
             "task", "summary", "subject", "проект", "работа"],
    "number": ["№", "номер", "код", "number", "no", "n", "wbs", "id задачи"],
    "key": ["id", "ключ", "key", "uid", "guid", "идентификатор"],
    "parent": ["родитель", "parent", "родительская задача", "parent id", "родитель id"],
    "level": ["уровень", "level", "outline level", "уровень структуры"],
    "executors": ["исполнитель", "исполнители", "ответственный", "ответственные", "текущий исполнитель",
                  "assignee", "assignees", "executor", "owner", "resource", "resources", "ресурсы"],
    "initiator": ["инициатор", "автор", "заказчик", "reporter", "creator", "author", "requester"],
    "planStart": ["начало", "дата начала", "план начала", "план начало", "старт", "start", "begin",
                  "дата начала план", "plan start", "start date", "planned start"],
    "planEnd": ["окончание", "срок", "дата окончания", "план окончания", "план окончание", "конец", "due",
                "due date", "finish", "end", "end date", "deadline", "крайняя дата", "крайняя дата выполнения",
                "выполнить до", "дата окончания план", "plan finish", "planned finish"],
    "factStart": ["факт начала", "фактическое начало", "факт начало", "actual start"],
    "factEnd": ["факт окончания", "фактическое окончание", "факт окончание", "дата выполнения", "дата закрытия",
                "closed", "resolved", "actual finish", "completed", "done date"],
    "percent": ["%", "процент", "выполнение", "прогресс", "progress", "percent", "complete", "% выполнения",
                "percent complete", "готовность"],
    "status": ["статус", "состояние", "status", "state", "статус задачи", "статус задания"],
    "priority": ["приоритет", "priority"],
    "urgency": ["срочность", "urgency"],
    "impact": ["влияние", "impact"],
    "category": ["категория", "тип", "category", "type", "вид", "вид задания"],
    "color": ["цвет", "color", "colour"],
}
STATUS_WORDS = [  # (вид, признаки в тексте статуса) — порядок важен: «не начат» раньше «начат»
    ("cancelled", ["отмен", "отклон", "cancel", "reject", "won't", "wont"]),
    ("notStarted", ["не начат", "не начата", "открыт", "новая", "новый", "new", "open", "to do", "todo", "backlog",
                    "очеред", "запланир", "planned"]),
    ("done", ["выполн", "заверш", "закрыт", "готов", "done", "closed", "resolved", "complete", "✔", "сделан"]),
    ("overdue", ["просроч", "overdue", "⚠"]),
    ("onHold", ["ожид", "отлож", "приостан", "hold", "waiting", "pause", "блокир", "blocked"]),
    ("release", ["релиз", "провер", "тест", "review", "testing", "qa", "согласов"]),
    ("inProgress", ["работ", "progress", "процесс", "▶", "active", "doing", "начат", "выполняется"]),
]
SCALE_WORDS = [(1, ["очень низ", "минимал", "lowest", "trivial"]), (2, ["низ", "low", "minor"]),
               (3, ["средн", "обычн", "нормал", "medium", "normal", "major"]),
               (5, ["очень выс", "критич", "блокер", "highest", "critical", "blocker", "срочн", "urgent"]),
               (4, ["выс", "high"])]
DEFAULT_LEVELS = {1: ["Задача"], 2: ["Проект", "Задача"], 3: ["Раздел", "Проект", "Задача"],
                  4: ["Раздел", "Группа", "Проект", "Задача"]}


@dataclass
class TableData:
    name: str
    rows: list[list]
    levels: list[int] = field(default_factory=list)     # группировка строк Excel
    indents: list[int] = field(default_factory=list)    # отступ текста в первой колонке
    sheets: list[str] = field(default_factory=list)
    sheet: str = ""
    path: str = ""

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


@dataclass
class Mapping:
    header_row: int = 0
    columns: dict = field(default_factory=dict)          # поле → номер колонки
    custom: dict = field(default_factory=dict)           # название пользовательского поля → номер колонки
    hierarchy: str = "flat"                              # outline | level | numbering | parent | indent | flat
    level_names: list = field(default_factory=list)
    status_mode: str = "dates"                           # column | dates
    status_map: dict = field(default_factory=dict)       # значение → вид статуса
    date_order: str = "dmy"                              # dmy | mdy | ymd
    template: str = ""
    signature: list = field(default_factory=list)
    sheet: str = ""

    def to_json(self) -> dict:
        return asdict(self)

    @classmethod
    def from_json(cls, d: dict) -> "Mapping":
        m = cls()
        for k, v in d.items():
            if hasattr(m, k):
                setattr(m, k, v)
        m.columns = {k: int(v) for k, v in m.columns.items()}
        m.custom = {k: int(v) for k, v in m.custom.items()}
        return m


# ---------- чтение ----------
def read_xlsx(path: str | Path, sheet: str | None = None) -> TableData:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    names = [ws.title for ws in wb.worksheets if not ws.title.startswith("_")]
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[names[0]]
    rows, levels, indents = [], [], []
    for r in ws.iter_rows():
        vals = [c.value for c in r]
        while vals and vals[-1] is None:
            vals.pop()
        rows.append(vals)
        idx = r[0].row
        levels.append(ws.row_dimensions[idx].outline_level or 0)
        first = next((c for c in r if c.value not in (None, "")), None)
        indents.append(int(first.alignment.indent) if first is not None and first.alignment else 0)
    return TableData(Path(path).name, rows, levels, indents, names, ws.title, str(path))


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def parse_text_table(text: str) -> list[list]:
    sample = text[:4096]
    delim = "\t" if "\t" in sample else None
    if delim is None:
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=";,|\t").delimiter
        except csv.Error:
            delim = ";" if sample.count(";") >= sample.count(",") else ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
    return [[c if c != "" else None for c in r] for r in rows]


def read_csv(path: str | Path) -> TableData:
    text = _decode(Path(path).read_bytes())
    rows = parse_text_table(text)
    return TableData(Path(path).name, rows, [0] * len(rows), _text_indents(rows), path=str(path))


def read_clipboard_text(text: str) -> TableData:
    rows = parse_text_table(text)
    return TableData("Таблица из буфера обмена", rows, [0] * len(rows), _text_indents(rows))


def _text_indents(rows: list[list]) -> list[int]:
    res = []
    for r in rows:
        first = next((c for c in r if c not in (None, "")), "")
        s = str(first)
        res.append((len(s) - len(s.lstrip(" "))) // 2)
    return res


# ---------- угадывание ----------
def norm(s) -> str:
    return re.sub(r"[\s_:]+", " ", str(s or "")).strip().lower().replace("ё", "е")


def signature(table: TableData, header_row: int) -> list[str]:
    return [norm(x) for x in (table.rows[header_row] if header_row < len(table.rows) else []) if x not in (None, "")]


def guess_header_row(table: TableData) -> int:
    best, best_score = 0, -1
    for i, r in enumerate(table.rows[:25]):
        texts = [norm(x) for x in r if isinstance(x, str) and x.strip()]
        score = sum(1 for t in texts for syns in SYNONYMS.values() if t in syns)
        if score > best_score and len(texts) >= 2:
            best, best_score = i, score
    return best


def guess_columns(headers: list) -> dict:
    cols: dict[str, int] = {}
    used: set[int] = set()
    normed = [norm(h) for h in headers]
    for exact in (True, False):
        for fld, syns in SYNONYMS.items():
            if fld in cols:
                continue
            for i, h in enumerate(normed):
                if i in used or not h:
                    continue
                if (h in syns) if exact else any(h.startswith(s) for s in syns if len(s) > 2):
                    cols[fld] = i
                    used.add(i)
                    break
    return cols


def guess_status_kind(value: str) -> str:
    v = norm(value)
    for kind, words in STATUS_WORDS:
        if any(w in v for w in words):
            return kind
    return "inProgress"


def guess_mapping(table: TableData) -> Mapping:
    m = Mapping(sheet=table.sheet)
    m.header_row = guess_header_row(table)
    headers = table.rows[m.header_row] if table.rows else []
    m.columns = guess_columns(headers)
    m.signature = signature(table, m.header_row)
    data = range(m.header_row + 1, len(table.rows))
    if any(table.levels[i] for i in data if i < len(table.levels)):
        m.hierarchy = "outline"
    elif "level" in m.columns:
        m.hierarchy = "level"
    elif "parent" in m.columns:
        m.hierarchy = "parent"
    elif "number" in m.columns and _dotted_share(table, m) > 0.5:
        m.hierarchy = "numbering"
    elif len({table.indents[i] for i in data if i < len(table.indents)}) > 1:
        m.hierarchy = "indent"
    else:
        m.hierarchy = "flat"
    if "key" not in m.columns and "number" in m.columns:  # уникальные номера — ещё и ключ строки
        col = m.columns["number"]
        vals = [r[col] for r in table.rows[m.header_row + 1:] if col < len(r) and r[col] not in (None, "")]
        if vals and len(set(map(str, vals))) == len(vals) and len(vals) >= 0.9 * len(table.rows[m.header_row + 1:]):
            m.columns["key"] = col
    if "status" in m.columns:
        m.status_mode = "column"
        m.status_map = {v: guess_status_kind(v) for v in distinct(table, m, m.columns["status"])}
    m.date_order = _guess_date_order(table, m)
    return m


def _dotted_share(table: TableData, m: Mapping) -> float:
    col = m.columns["number"]
    vals = [str(r[col]) for r in table.rows[m.header_row + 1:] if col < len(r) and r[col] not in (None, "")]
    return sum(1 for v in vals if re.fullmatch(r"\d+(\.\d+)*\.?", v.strip())) / max(1, len(vals)) if vals else 0


def _guess_date_order(table: TableData, m: Mapping) -> str:
    for fld in ("planStart", "planEnd", "factStart", "factEnd"):
        col = m.columns.get(fld)
        if col is None:
            continue
        for r in table.rows[m.header_row + 1:]:
            v = r[col] if col < len(r) else None
            if isinstance(v, str):
                mm = re.match(r"\s*(\d{1,2})[./](\d{1,2})[./]\d{2,4}", v)
                if mm and int(mm[1]) <= 12 < int(mm[2]):
                    return "mdy"
                if re.match(r"\s*\d{4}-\d{2}-\d{2}", v):
                    return "ymd"
    return "dmy"


def distinct(table: TableData, m: Mapping, col: int) -> list[str]:
    seen = []
    for r in table.rows[m.header_row + 1:]:
        v = r[col] if col < len(r) else None
        if v not in (None, "") and str(v).strip() not in seen:
            seen.append(str(v).strip())
    return seen


# ---------- разбор значений ----------
def parse_date(v, order: str = "dmy") -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:  # число-дата Excel
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    if s in ("", "—", "-", "–"):
        return None
    s = s.split(" ")[0].split("T")[0]
    fmts = {"dmy": ("%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"),
            "mdy": ("%m/%d/%Y", "%m/%d/%y", "%m.%d.%Y", "%Y-%m-%d"),
            "ymd": ("%Y-%m-%d", "%Y.%m.%d", "%Y/%m/%d", "%d.%m.%Y")}[order]
    for f in fmts:
        try:
            return datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    return None


def parse_percent(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(v * 100 if 0 < v <= 1 and not float(v).is_integer() else v, 1)
    mm = re.search(r"(\d+(?:[.,]\d+)?)\s*%", str(v)) or re.search(r"(\d+(?:[.,]\d+)?)", str(v))
    if not mm:
        return None
    x = float(mm[1].replace(",", "."))
    return round(x * 100, 1) if x <= 1 and "%" not in str(v) and "." in mm[1].replace(",", ".") else min(x, 100)


def split_people(v) -> list[str]:
    if v in (None, ""):
        return []
    return [p.strip() for p in re.split(r"[,;\n]+", str(v)) if p.strip()]


def short_name(full: str) -> str:
    p = full.split()
    return f"{p[0]} {p[1][0]}.{p[2][0]}." if len(p) == 3 and all(p) else full


def scale_order(value: str, fallback: int) -> int:
    v = norm(value)
    for order, words in SCALE_WORDS:
        if any(w in v for w in words):
            return order
    return fallback


# ---------- сборка документа ----------
def convert(table: TableData, m: Mapping, source: str | None = None) -> tuple[dict, list[str]]:
    """Строки таблицы → документ gantt-exchange (kind = full) и предупреждения."""
    warnings: list[str] = []
    cols = m.columns

    def cell(r, fld):
        c = cols.get(fld)
        return r[c] if c is not None and c < len(r) else None

    rows = []
    for i in range(m.header_row + 1, len(table.rows)):
        r = table.rows[i]
        name = cell(r, "name")
        if name in (None, "") or not str(name).strip():
            continue
        rows.append((i, r, str(name).strip()))
    if not rows:
        raise ValueError("В таблице нет строк с наименованием. Проверьте, какая колонка — «Наименование».")

    # глубина строки
    depths: list[int] = []
    if m.hierarchy == "outline":
        base = min(table.levels[i] for i, _r, _n in rows)
        depths = [table.levels[i] - base for i, _r, _n in rows]
    elif m.hierarchy == "indent":
        found = sorted({table.indents[i] for i, _r, _n in rows})
        depths = [found.index(table.indents[i]) for i, _r, _n in rows]
    elif m.hierarchy == "numbering":
        depths = [max(0, str(cell(r, "number") or "").strip().rstrip(".").count(".")) for _i, r, _n in rows]
    elif m.hierarchy == "level":
        raw = [cell(r, "level") for _i, r, _n in rows]
        if all(isinstance(x, (int, float)) or (isinstance(x, str) and x.strip().isdigit()) for x in raw if x):
            nums = [int(float(x)) if x not in (None, "") else 1 for x in raw]
            lo = min(nums)
            depths = [n - lo for n in nums]
        else:
            order = []
            for x in raw:
                if x and str(x) not in order:
                    order.append(str(x))
            depths = [order.index(str(x)) if x else 0 for x in raw]
            if not m.level_names:
                m.level_names = order
    else:
        depths = [0] * len(rows)

    max_depth = max(depths) + 1
    names = list(m.level_names) if len(m.level_names) >= max_depth else \
        DEFAULT_LEVELS.get(max_depth, [f"Уровень {k + 1}" for k in range(max_depth)])
    levels = [{"id": f"L{k + 1}", "name": names[k] if k < len(names) else f"Уровень {k + 1}", "order": k + 1}
              for k in range(max_depth)]

    people: dict[str, dict] = {}

    def person(nm: str) -> str:
        pid = str(uuid.uuid5(NS, "person|" + nm))
        people.setdefault(pid, {"id": pid, "name": short_name(nm), "fullName": nm})
        return pid

    scales = {"priority": {}, "urgency": {}, "impact": {}, "category": {}}
    statuses: dict[str, dict] = {}
    custom_values: dict[str, list] = {k: [] for k in m.custom}
    nodes, keys, stack, paths = [], {}, [], {}
    for pos, ((i, r, name), depth) in enumerate(zip(rows, depths)):
        stack = stack[:depth]
        parent_node = stack[-1] if stack else None
        key = cell(r, "key")
        path = (paths.get(parent_node["id"], "") if parent_node else "") + "/" + name
        occurrence = sum(1 for p in paths.values() if p == path)
        nid = str(key).strip() if key not in (None, "") else str(uuid.uuid5(NS, f"row|{path}|{occurrence}"))
        if nid in keys:
            nid = f"{nid}#{pos}"
            warnings.append(f"Ключ «{key}» повторяется — у повтора добавлен номер строки")
        node = {"id": nid, "parentId": parent_node["id"] if parent_node else None, "level": f"L{depth + 1}",
                "order": pos, "name": name}
        paths[nid] = path
        num = cell(r, "number")
        if num not in (None, ""):
            node["number"] = str(num).strip()
        for fld in ("planStart", "planEnd", "factStart", "factEnd"):
            d = parse_date(cell(r, fld), m.date_order)
            if d:
                node[fld] = d
            elif cell(r, fld) not in (None, "", "—", "-"):
                warnings.append(f"Строка {i + 1}: не удалось прочитать дату «{cell(r, fld)}» ({FIELD_TITLES[fld]})")
        ex = [person(p) for p in split_people(cell(r, "executors"))]
        if ex:
            node["executors"] = ex
        ini = split_people(cell(r, "initiator"))
        if ini:
            node["initiator"] = person(ini[0])
        pct = parse_percent(cell(r, "percent"))
        if pct is not None:
            node["percent"] = pct
        st = cell(r, "status")
        if m.status_mode == "column" and st not in (None, ""):
            sv = str(st).strip()
            sid = "st:" + sv
            statuses.setdefault(sid, {"id": sid, "name": sv, "kind": m.status_map.get(sv) or guess_status_kind(sv),
                                      "order": len(statuses) + 1})
            node["status"] = sid
        for fld in scales:
            v = cell(r, fld)
            if v not in (None, ""):
                sv = str(v).strip()
                scales[fld].setdefault(sv, len(scales[fld]) + 1)
                node[fld] = f"{fld[:3]}:{sv}"
        color = cell(r, "color")
        if isinstance(color, str) and re.fullmatch(r"#[0-9A-Fa-f]{6}", color.strip()):
            node["color"] = color.strip().upper()
        for cname, ccol in m.custom.items():
            v = r[ccol] if ccol < len(r) else None
            if v not in (None, ""):
                custom_values[cname].append(v)
                node.setdefault("custom", {})[_cid(cname)] = v if isinstance(v, (int, float)) else str(v)
        keys[nid] = node
        nodes.append(node)
        stack.append(node)

    if m.hierarchy == "parent":  # связь по колонке «Родитель»
        for (i, r, _n), node in zip(rows, nodes):
            p = cell(r, "parent")
            node["parentId"] = str(p).strip() if p not in (None, "") and str(p).strip() in keys else None
        _depths_from_parents(nodes, levels, m)

    if m.status_mode == "dates" or not statuses:
        statuses = {"inProgress": {"id": "inProgress", "name": "В работе", "kind": "inProgress", "order": 1},
                    "done": {"id": "done", "name": "Выполнено", "kind": "done", "order": 2}}
        leaf_order = max(x["order"] for x in levels)
        for n in nodes:
            n.pop("status", None)
            if int(n["level"][1:]) == leaf_order:
                n["status"] = "done" if n.get("factEnd") or (n.get("percent") or 0) >= 100 else "inProgress"

    dictionaries = {"levels": levels, "statuses": list(statuses.values()),
                    "people": sorted(people.values(), key=lambda p: p["name"])}
    keymap = {"priority": "priorities", "urgency": "urgencies", "impact": "impacts", "category": "categories"}
    for fld, vals in scales.items():
        if vals:
            items = [{"id": f"{fld[:3]}:{v}", "name": v} for v in vals]
            if fld != "category":
                for it, (v, pos) in zip(items, vals.items()):
                    it["order"] = scale_order(v, pos)
            dictionaries[keymap[fld]] = items
    if m.custom:
        fields = []
        for cname, vals in custom_values.items():
            numeric = vals and all(isinstance(v, (int, float)) for v in vals)
            f = {"id": _cid(cname), "name": cname, "type": "number" if numeric else "string"}
            fields.append(f)
        dictionaries["customFields"] = fields
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    doc = {"format": "gantt-exchange", "schemaVersion": "1.0",
           "meta": {"source": source or f"table:{m.template or 'Таблица'}",
                    "sourceName": m.template or table.name, "title": table.name, "exportedAt": stamp,
                    "kind": "full", "cursor": stamp},
           "dictionaries": dictionaries, "nodes": nodes}
    return doc, warnings


def _cid(name: str) -> str:
    """id пользовательского поля: латиница и цифры (схема), стабильно от названия."""
    base = re.sub(r"[^A-Za-z0-9_]", "", name)
    return ("f" + base if not base[:1].isalpha() else base)[:20] + "_" + uuid.uuid5(NS, name).hex[:8]


def _depths_from_parents(nodes: list[dict], levels: list[dict], m: Mapping) -> None:
    by_id = {n["id"]: n for n in nodes}

    def depth(n, seen=()):
        p = n.get("parentId")
        if not p or p not in by_id or p in seen:
            return 0
        return depth(by_id[p], seen + (n["id"],)) + 1

    ds = [depth(n) for n in nodes]
    max_depth = max(ds) + 1
    names = list(m.level_names) if len(m.level_names) >= max_depth else \
        DEFAULT_LEVELS.get(max_depth, [f"Уровень {k + 1}" for k in range(max_depth)])
    levels[:] = [{"id": f"L{k + 1}", "name": names[k] if k < len(names) else f"Уровень {k + 1}", "order": k + 1}
                 for k in range(max_depth)]
    for n, d in zip(nodes, ds):
        n["level"] = f"L{d + 1}"
